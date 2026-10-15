#!/usr/bin/env python3
"""Export data STO (hasil filter) ke xlsx.
Input : JSON stdin {"out": path, "rows": [...], "meta": {...}}
Output: JSON stdout {"ok": true, "rows": N}
"""
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ('code', 'Kode', None),
    ('name', 'Nama Item', None),
    ('group_name', 'Group', None),
    ('location_name', 'Lokasi', None),
    ('device_name', 'Device', None),
    ('unit', 'Satuan', None),
    ('expected_qty', 'Stok Sistem', '#,##0'),
    ('counted_qty', 'Hitung', '#,##0'),
    ('count_status', 'Status Hitung', None),
    ('audit_status', 'Status Audit', None),
]

def main():
    data = json.load(sys.stdin)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data STO'
    headers = [h for _, h, _ in COLUMNS]
    ws.append(headers)
    head_fill = PatternFill('solid', fgColor='1E293B')
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = head_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = 18
    for r in data.get('rows', []):
        ws.append([r.get(k) if r.get(k) is not None else '' for k, _, _ in COLUMNS])
    for idx, (_, _, fmt) in enumerate(COLUMNS, start=1):
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                for cell in row:
                    cell.number_format = fmt
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    meta = wb.create_sheet('Info')
    meta.append(['Keterangan', 'Nilai'])
    for k, v in (data.get('meta') or {}).items():
        meta.append([k, str(v)])
    meta.column_dimensions['A'].width = 20
    meta.column_dimensions['B'].width = 40
    wb.save(data['out'])
    print(json.dumps({'ok': True, 'rows': len(data.get('rows', []))}))

if __name__ == '__main__':
    try:
        main()
    except Exception as e:  # noqa: BLE001
        print(json.dumps({'ok': False, 'error': str(e)}))
        sys.exit(1)
